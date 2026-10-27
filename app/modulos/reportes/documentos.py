"""Documentos de las actas (HU-11): el PDF de cada acta y los reportes en PDF y Excel.

Funciones puras: reciben actas ya cargadas (con sus relaciones) y devuelven bytes.

El PDF de un acta es **determinista** (``invariant=True``): el mismo contenido produce
exactamente los mismos bytes, sin fecha de creación ni identificador aleatorio. Por eso su
SHA-256 sirve como huella de la conformidad: si alguien altera el documento, la huella deja
de coincidir.
"""

import hashlib
from collections.abc import Sequence
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.comun.fechas import ZONA_COLOMBIA
from app.modulos.reportes.modelos import Acta, TipoActa

TITULOS = {TipoActa.ENTREGA: "Acta de entrega", TipoActa.RECEPCION: "Acta de recepción"}
_FORMATO_FECHA = "%d/%m/%Y %H:%M"


def _fecha(acta: Acta) -> str:
    return acta.creado_en.astimezone(ZONA_COLOMBIA).strftime(_FORMATO_FECHA)


def operacion(acta: Acta) -> str:
    """ "Alquiler #12 (del 05/11/2026 al 07/11/2026)" o "Compra #8"."""
    if acta.alquiler is not None:
        a = acta.alquiler
        return f"Alquiler #{a.id} (del {a.fecha_inicio:%d/%m/%Y} al {a.fecha_fin:%d/%m/%Y})"
    return f"Compra #{acta.venta_id}"


def cliente(acta: Acta):
    return acta.alquiler.cliente if acta.alquiler else acta.venta.cliente


def _bateria(acta: Acta) -> str:
    return "No aplica" if acta.nivel_bateria is None else f"{acta.nivel_bateria} %"


def _documento(salida: BytesIO, titulo: str, **opciones) -> SimpleDocTemplate:
    return SimpleDocTemplate(
        salida,
        title=titulo,
        author="TransportExperience",
        leftMargin=2 * cm,
        rightMargin=2 * cm,
        topMargin=2 * cm,
        bottomMargin=2 * cm,
        invariant=True,  # mismos datos → mismos bytes (ver el docstring del módulo)
        **opciones,
    )


def pdf_acta(acta: Acta) -> bytes:
    """El acta en PDF (tamaño carta). No incluye la conformidad: es el documento que se acepta."""
    salida = BytesIO()
    estilos = getSampleStyleSheet()
    titulo = f"{TITULOS[acta.tipo]} {acta.numero}"
    usuario_cliente = cliente(acta)
    filas = [
        ["Fecha", _fecha(acta)],
        ["Operación", operacion(acta)],
        ["Entrega a domicilio", f"Domicilio #{acta.domicilio_id}" if acta.domicilio_id else "No"],
        ["Cliente", f"{usuario_cliente.nombre_completo} ({usuario_cliente.correo})"],
        ["Vehículo", f"{acta.vehiculo.codigo} — {acta.vehiculo.modelo.nombre}"],
        ["Batería", _bateria(acta)],
        ["Estado del vehículo", Paragraph(acta.estado_vehiculo, estilos["Normal"])],
        ["Registrada por", acta.registrada_por.nombre_completo],
    ]
    tabla = Table(filas, colWidths=[5 * cm, 12 * cm])
    tabla.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.lightgrey),
            ]
        )
    )
    contenido = [
        Paragraph("TransportExperience", estilos["Title"]),
        Paragraph(titulo, estilos["Heading2"]),
        Spacer(1, 0.4 * cm),
        tabla,
        Spacer(1, 0.8 * cm),
        Paragraph(
            "Con la conformidad del cliente en la plataforma, queda registrado quién la dio, "
            "cuándo y la huella SHA-256 de este documento.",
            estilos["Italic"],
        ),
    ]
    _documento(salida, titulo, pagesize=letter).build(contenido)
    return salida.getvalue()


def huella(pdf: bytes) -> str:
    """SHA-256 en hexadecimal (64 caracteres), el formato de ``actas.hash_pdf``."""
    return hashlib.sha256(pdf).hexdigest()


COLUMNAS = [
    "Número",
    "Tipo",
    "Fecha",
    "Operación",
    "Domicilio",
    "Cliente",
    "Vehículo",
    "Modelo",
    "Batería",
    "Estado del vehículo",
    "Registrada por",
    "Conformidad",
]


def _fila(acta: Acta) -> list:
    conformidad = (
        f"{acta.conforme_por.nombre_completo}, "
        f"{acta.conforme_en.astimezone(ZONA_COLOMBIA).strftime(_FORMATO_FECHA)}"
        if acta.conforme_en
        else "Pendiente"
    )
    return [
        acta.numero,
        "Entrega" if acta.tipo == TipoActa.ENTREGA else "Recepción",
        _fecha(acta),
        operacion(acta),
        f"#{acta.domicilio_id}" if acta.domicilio_id else "",
        cliente(acta).nombre_completo,
        acta.vehiculo.codigo,
        acta.vehiculo.modelo.nombre,
        _bateria(acta),
        acta.estado_vehiculo,
        acta.registrada_por.nombre_completo,
        conformidad,
    ]


def excel_reporte(actas: Sequence[Acta], titulo: str) -> bytes:
    """Una hoja con una fila por acta, lista para filtrar en Excel."""
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Actas"
    hoja.append([titulo])
    hoja["A1"].font = Font(bold=True, size=13)
    hoja.append([])
    hoja.append(COLUMNAS)
    for celda in hoja[3]:
        celda.font = Font(bold=True)
    for acta in actas:
        hoja.append(_fila(acta))
    hoja.auto_filter.ref = (
        f"A3:{hoja.cell(row=3, column=len(COLUMNAS)).column_letter}{hoja.max_row}"
    )
    hoja.freeze_panes = "A4"
    anchos = [12, 11, 17, 40, 10, 28, 12, 26, 9, 50, 26, 30]
    for indice, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[hoja.cell(row=3, column=indice).column_letter].width = ancho
    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()


def pdf_reporte(actas: Sequence[Acta], titulo: str) -> bytes:
    """Las actas del rango en una tabla, en horizontal."""
    salida = BytesIO()
    estilos = getSampleStyleSheet()
    chica = estilos["BodyText"].clone("chica", fontSize=7, leading=8.5)
    columnas = [
        "Número",
        "Tipo",
        "Fecha",
        "Operación",
        "Cliente",
        "Vehículo",
        "Estado",
        "Conformidad",
    ]
    filas = [columnas]
    for acta in actas:
        f = _fila(acta)
        datos = [f[0], f[1], f[2], f[3], f[5], f"{f[6]} — {f[7]}", f[9], f[11]]
        filas.append([Paragraph(str(valor), chica) for valor in datos])
    if len(filas) == 1:
        filas.append(["Sin actas en el rango."] + [""] * (len(columnas) - 1))
    tabla = Table(
        filas,
        colWidths=[2 * cm, 1.8 * cm, 2.5 * cm, 4.5 * cm, 3.5 * cm, 3.8 * cm, 4.6 * cm, 3 * cm],
        repeatRows=1,
    )
    tabla.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, 0), 8),
                ("LINEBELOW", (0, 0), (-1, 0), 0.8, colors.black),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )
    contenido = [
        Paragraph("TransportExperience", estilos["Title"]),
        Paragraph(titulo, estilos["Heading2"]),
        Paragraph(f"{len(actas)} acta(s).", estilos["Normal"]),
        Spacer(1, 0.4 * cm),
        tabla,
    ]
    _documento(salida, titulo, pagesize=landscape(letter)).build(contenido)
    return salida.getvalue()
