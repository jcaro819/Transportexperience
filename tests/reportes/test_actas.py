"""Actas de entrega y recepción (HU-11): se generan solas, PDF determinista, conformidad, reportes.

Contra Postgres, con "hoy" fijo.
"""

from datetime import UTC, date, datetime, timedelta
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.comun.errores import PermisoDenegado
from app.comun.paginacion import ParametrosPagina
from app.modulos.alquileres import servicio as alquileres
from app.modulos.alquileres.excepciones import FueraDeLasFechasDelAlquiler
from app.modulos.domicilios import servicio as domicilios
from app.modulos.domicilios.modelos import EstadoDomicilio
from app.modulos.inventario.excepciones import BateriaNoAplica
from app.modulos.pagos import procesamiento
from app.modulos.pagos.esquemas import PagoCrear
from app.modulos.pagos.modelos import MetodoPago
from app.modulos.pagos.pasarela import PasarelaFalsa
from app.modulos.reportes import servicio
from app.modulos.reportes.documentos import huella, pdf_acta
from app.modulos.reportes.excepciones import (
    ActaNoEncontrada,
    ConformidadYaRegistrada,
    RangoDeReporteInvalido,
)
from app.modulos.reportes.modelos import Acta, TipoActa
from app.modulos.reportes.registro import SIN_OBSERVACIONES, DatosActa
from app.modulos.usuarios.modelos import Rol
from app.modulos.ventas import servicio as ventas
from app.modulos.ventas.esquemas import VentaCrear

HOY = date(2026, 11, 2)
MOMENTO = datetime(2026, 11, 2, 15, 0, tzinfo=UTC)


@pytest.fixture(autouse=True)
def metodo(sesion) -> None:
    sesion.add(MetodoPago(codigo="tarjeta", nombre="Tarjeta"))
    sesion.flush()


@pytest.fixture
def cliente(fabrica):
    return fabrica.cliente()


@pytest.fixture
def operador(crear_usuario):
    return crear_usuario(Rol.OPERADOR)


def _pagar(sesion, cliente, **operacion):
    procesamiento.pagar(
        sesion,
        cliente,
        PagoCrear(metodo_pago="tarjeta"),
        PasarelaFalsa(),
        momento=MOMENTO,
        **operacion,
    )


@pytest.fixture
def alquiler(sesion, fabrica, cliente):
    """Un alquiler pagado que empieza hoy, de un monopatín."""
    modelo = fabrica.modelo(nombre="Monopatín Urbano 350", tarifa_diaria=30_000)
    fabrica.vehiculo(modelo, codigo="MON-0001", nivel_bateria=90)
    alquiler = alquileres.crear_alquiler(
        sesion, cliente, modelo.id, HOY, HOY + timedelta(days=1), hoy=HOY, momento=MOMENTO
    )
    _pagar(sesion, cliente, alquiler_id=alquiler.id)
    return alquiler


def _actas(sesion) -> list[Acta]:
    return list(sesion.scalars(select(Acta).order_by(Acta.id)))


def _entregar(sesion, alquiler, operador, datos=None):
    return alquileres.entregar_vehiculo(sesion, alquiler.id, operador, datos, hoy=HOY)


# --- Se generan solas (decisión 2) ---------------------------------------------------------


def test_la_entrega_en_el_local_genera_el_acta_de_entrega(sesion, alquiler, operador) -> None:
    _entregar(sesion, alquiler, operador, DatosActa("Sin rayones, frenos bien.", 88))

    [acta] = _actas(sesion)
    assert acta.tipo == TipoActa.ENTREGA
    assert (acta.alquiler_id, acta.vehiculo_id, acta.domicilio_id) == (
        alquiler.id,
        alquiler.vehiculo_id,
        None,
    )
    assert (acta.estado_vehiculo, acta.nivel_bateria) == ("Sin rayones, frenos bien.", 88)
    assert acta.registrada_por_id == operador.id
    assert acta.conforme_en is None


def test_la_recepcion_genera_el_acta_de_recepcion(sesion, alquiler, operador) -> None:
    _entregar(sesion, alquiler, operador)
    alquileres.recibir_vehiculo(
        sesion, alquiler.id, operador, DatosActa("Rayón en el guardabarros.")
    )

    entrega, recepcion = _actas(sesion)
    assert (entrega.tipo, entrega.estado_vehiculo) == (TipoActa.ENTREGA, SIN_OBSERVACIONES)
    assert (recepcion.tipo, recepcion.estado_vehiculo) == (
        TipoActa.RECEPCION,
        "Rayón en el guardabarros.",
    )


def test_si_la_entrega_falla_no_queda_acta(sesion, alquiler, operador) -> None:
    with pytest.raises(FueraDeLasFechasDelAlquiler):
        alquileres.entregar_vehiculo(sesion, alquiler.id, operador, hoy=HOY - timedelta(days=1))
    assert _actas(sesion) == []


def test_una_cicla_no_lleva_bateria_en_el_acta(sesion, fabrica, cliente, operador) -> None:
    cicla = fabrica.modelo(fabrica.tipo("Cicla", usa_bateria=False), tarifa_diaria=20_000)
    fabrica.vehiculo(cicla)
    alquiler = alquileres.crear_alquiler(
        sesion, cliente, cicla.id, HOY, HOY, hoy=HOY, momento=MOMENTO
    )
    _pagar(sesion, cliente, alquiler_id=alquiler.id)

    with pytest.raises(BateriaNoAplica):
        _entregar(sesion, alquiler, operador, DatosActa("Todo bien.", 50))
    assert _actas(sesion) == []


def _solicitar_domicilio(sesion, fabrica, cliente, *, alquiler=None, venta=None):
    fabrica.zona()
    domicilio = domicilios.solicitar_domicilio(
        sesion,
        cliente,
        alquiler_id=alquiler.id if alquiler else None,
        venta_id=venta.id if venta else None,
        direccion="Calle 36 # 19-20",
        indicaciones=None,
        latitud=Decimal("7.119300"),
        longitud=Decimal("-73.122700"),
        momento=MOMENTO,
    ).domicilio
    return domicilio


def test_la_entrega_por_domicilio_genera_el_acta(
    sesion, fabrica, cliente, operador, crear_usuario
) -> None:
    modelo = fabrica.modelo(tarifa_diaria=30_000)
    fabrica.vehiculo(modelo)
    alquiler = alquileres.crear_alquiler(
        sesion, cliente, modelo.id, HOY, HOY, hoy=HOY, momento=MOMENTO
    )
    domicilio = _solicitar_domicilio(sesion, fabrica, cliente, alquiler=alquiler)
    _pagar(sesion, cliente, alquiler_id=alquiler.id)
    domiciliario = crear_usuario(Rol.DOMICILIARIO)
    domicilios.asignar_domiciliario(sesion, domicilio.id, operador, domiciliario.id, hoy=HOY)
    for estado in (EstadoDomicilio.RECOGIDO, EstadoDomicilio.EN_CAMINO):
        domicilios.avanzar_estado(sesion, domicilio.id, domiciliario, estado, hoy=HOY)

    domicilios.avanzar_estado(
        sesion,
        domicilio.id,
        domiciliario,
        EstadoDomicilio.ENTREGADO,
        DatosActa("Entregado en portería."),
        hoy=HOY,
    )

    [acta] = _actas(sesion)
    assert (acta.tipo, acta.domicilio_id, acta.alquiler_id) == (
        TipoActa.ENTREGA,
        domicilio.id,
        alquiler.id,
    )
    assert acta.registrada_por_id == domiciliario.id


def test_una_compra_entregada_lleva_un_acta_por_unidad(
    sesion, fabrica, cliente, operador, crear_usuario
) -> None:
    modelo = fabrica.modelo(precio_venta=900_000)
    fabrica.vehiculo(modelo)
    fabrica.vehiculo(modelo)
    casco = fabrica.producto()
    venta = ventas.crear_venta(
        sesion,
        cliente,
        VentaCrear(
            vehiculos=[{"modelo_id": modelo.id, "cantidad": 2}],
            productos=[{"producto_id": casco.id}],
        ),
        momento=MOMENTO,
    )
    domicilio = _solicitar_domicilio(sesion, fabrica, cliente, venta=venta)
    _pagar(sesion, cliente, venta_id=venta.id)
    domiciliario = crear_usuario(Rol.DOMICILIARIO)
    domicilios.asignar_domiciliario(sesion, domicilio.id, operador, domiciliario.id, hoy=HOY)
    for estado in (EstadoDomicilio.RECOGIDO, EstadoDomicilio.EN_CAMINO, EstadoDomicilio.ENTREGADO):
        domicilios.avanzar_estado(sesion, domicilio.id, domiciliario, estado, hoy=HOY)

    actas = _actas(sesion)
    assert len(actas) == 2  # el casco no es un vehículo: no lleva acta
    assert {a.venta_id for a in actas} == {venta.id}


# --- PDF determinista y conformidad --------------------------------------------------------


@pytest.fixture
def acta(sesion, alquiler, operador) -> Acta:
    _entregar(sesion, alquiler, operador, DatosActa("Sin rayones.", 88))
    [acta] = _actas(sesion)
    return acta


def test_el_pdf_del_acta_es_siempre_el_mismo(sesion, acta, cliente) -> None:
    primero = servicio.pdf(sesion, acta.id, cliente)
    segundo = servicio.pdf(sesion, acta.id, cliente)

    assert primero.startswith(b"%PDF")
    assert primero == segundo  # byte a byte: por eso su SHA-256 sirve de huella


def test_si_el_acta_cambia_cambia_la_huella(sesion, acta) -> None:
    antes = huella(pdf_acta(acta))
    acta.estado_vehiculo = "Otro texto."
    assert huella(pdf_acta(acta)) != antes


def test_el_cliente_da_su_conformidad_con_la_huella_del_pdf(sesion, acta, cliente) -> None:
    resultado = servicio.dar_conformidad(sesion, acta.id, cliente, momento=MOMENTO)

    assert resultado.conforme_por_id == cliente.id
    assert resultado.conforme_en == MOMENTO
    assert resultado.hash_pdf == huella(servicio.pdf(sesion, acta.id, cliente))
    assert len(resultado.hash_pdf) == 64
    assert servicio.verificar(sesion, acta.id, cliente) is True


def test_la_verificacion_detecta_un_acta_alterada(sesion, acta, cliente) -> None:
    servicio.dar_conformidad(sesion, acta.id, cliente)
    acta.estado_vehiculo = "Alguien cambió el texto después."
    sesion.flush()
    assert servicio.verificar(sesion, acta.id, cliente) is False


def test_sin_conformidad_no_hay_nada_que_verificar(sesion, acta, cliente) -> None:
    assert servicio.verificar(sesion, acta.id, cliente) is None


def test_la_conformidad_se_da_una_sola_vez(sesion, acta, cliente) -> None:
    servicio.dar_conformidad(sesion, acta.id, cliente)
    with pytest.raises(ConformidadYaRegistrada):
        servicio.dar_conformidad(sesion, acta.id, cliente)


def test_la_conformidad_es_del_cliente_de_la_operacion(
    sesion, acta, operador, crear_usuario
) -> None:
    with pytest.raises(PermisoDenegado):
        servicio.dar_conformidad(sesion, acta.id, operador)
    with pytest.raises(ActaNoEncontrada):
        servicio.dar_conformidad(sesion, acta.id, crear_usuario(Rol.CLIENTE))
    assert acta.conforme_en is None


# --- Listas y reportes ---------------------------------------------------------------------


def test_el_cliente_solo_ve_sus_actas(sesion, acta, cliente, crear_usuario, operador) -> None:
    otro = crear_usuario(Rol.CLIENTE)
    assert [a.id for a in servicio.listar_actas(sesion, ParametrosPagina(), cliente).elementos] == [
        acta.id
    ]
    assert servicio.listar_actas(sesion, ParametrosPagina(), otro).elementos == []
    assert servicio.listar_actas(sesion, ParametrosPagina(), operador).total == 1
    with pytest.raises(PermisoDenegado):
        servicio.listar_actas(sesion, ParametrosPagina(), crear_usuario(Rol.DOMICILIARIO))


def test_filtros_por_tipo_y_fechas(sesion, acta, operador) -> None:
    creada = acta.creado_en.date()

    def total(**filtros):
        return servicio.listar_actas(sesion, ParametrosPagina(), operador, **filtros).total

    assert total(tipo=TipoActa.ENTREGA) == 1
    assert total(tipo=TipoActa.RECEPCION) == 0
    assert total(desde=creada - timedelta(days=1), hasta=creada + timedelta(days=1)) == 1
    assert total(desde=creada + timedelta(days=2)) == 0


def test_reporte_en_excel(sesion, acta, operador) -> None:
    from app.modulos.reportes.documentos import COLUMNAS, excel_reporte

    datos = servicio.reporte(
        sesion, operador, desde=acta.creado_en.date(), hasta=acta.creado_en.date()
    )
    hoja = load_workbook(BytesIO(excel_reporte(datos.actas, datos.titulo))).active

    assert [c.value for c in hoja[3]] == COLUMNAS
    fila = [c.value for c in hoja[4]]
    assert fila[0] == acta.numero
    assert fila[6] == "MON-0001"
    assert fila[11] == "Pendiente"
    assert hoja.max_row == 4


def test_reporte_en_pdf(sesion, acta, operador) -> None:
    from app.modulos.reportes.documentos import pdf_reporte

    datos = servicio.reporte(sesion, operador)
    assert pdf_reporte(datos.actas, datos.titulo).startswith(b"%PDF")
    assert pdf_reporte([], "Vacío").startswith(b"%PDF")


@pytest.mark.parametrize(
    ("desde", "hasta"),
    [(date(2026, 11, 5), date(2026, 11, 1)), (date(2026, 1, 1), date(2026, 6, 30))],
    ids=["al_reves", "mas_de_92_dias"],
)
def test_rango_de_reporte_invalido(sesion, operador, desde, hasta) -> None:
    with pytest.raises(RangoDeReporteInvalido):
        servicio.reporte(sesion, operador, desde=desde, hasta=hasta)


def test_el_cliente_no_exporta_reportes(sesion, cliente) -> None:
    with pytest.raises(PermisoDenegado):
        servicio.reporte(sesion, cliente)
