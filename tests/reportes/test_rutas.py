"""Rutas de las actas (HU-11): descargas, conformidad por la API y permisos."""

import hashlib
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.comun import fechas
from app.modulos.pagos.modelos import MetodoPago
from app.modulos.usuarios.modelos import Rol
from tests.conftest import encabezado_de


@pytest.fixture
def entregado(sesion, fabrica, crear_usuario, cliente_api) -> dict:
    """Un alquiler entregado por la API, con su acta. Devuelve encabezados e ids."""
    sesion.add(MetodoPago(codigo="tarjeta", nombre="Tarjeta"))
    modelo = fabrica.modelo(tarifa_diaria=30_000)
    fabrica.vehiculo(modelo, codigo="MON-0001")
    cliente = encabezado_de(fabrica.cliente())
    operador = encabezado_de(crear_usuario(Rol.OPERADOR))
    hoy = str(fechas.hoy())
    alquiler = cliente_api.post(
        "/alquileres",
        json={"modelo_id": modelo.id, "fecha_inicio": hoy, "fecha_fin": hoy},
        headers=cliente,
    ).json()
    cliente_api.post(
        f"/alquileres/{alquiler['id']}/pago", json={"metodo_pago": "tarjeta"}, headers=cliente
    )
    entrega = cliente_api.post(
        f"/alquileres/{alquiler['id']}/entrega",
        json={"estado_vehiculo": "Sin rayones.", "nivel_bateria": 90},
        headers=operador,
    )
    assert entrega.status_code == 200
    [acta] = cliente_api.get("/actas/mias", headers=cliente).json()["elementos"]
    return {"cliente": cliente, "operador": operador, "acta": acta, "alquiler": alquiler}


def test_la_entrega_sin_cuerpo_sigue_funcionando(
    cliente_api, sesion, fabrica, crear_usuario
) -> None:
    sesion.add(MetodoPago(codigo="tarjeta", nombre="Tarjeta"))
    modelo = fabrica.modelo()
    fabrica.vehiculo(modelo)
    cliente = encabezado_de(fabrica.cliente())
    hoy = str(fechas.hoy())
    alquiler = cliente_api.post(
        "/alquileres",
        json={"modelo_id": modelo.id, "fecha_inicio": hoy, "fecha_fin": hoy},
        headers=cliente,
    ).json()
    cliente_api.post(
        f"/alquileres/{alquiler['id']}/pago", json={"metodo_pago": "tarjeta"}, headers=cliente
    )
    respuesta = cliente_api.post(
        f"/alquileres/{alquiler['id']}/entrega", headers=encabezado_de(crear_usuario(Rol.OPERADOR))
    )
    assert respuesta.status_code == 200  # el contrato del Sprint 2 no se rompe


def test_el_acta_trae_lo_anotado(entregado) -> None:
    acta = entregado["acta"]
    assert acta["tipo"] == "entrega"
    assert acta["numero"].startswith("ACT-")
    assert (acta["estado_vehiculo"], acta["nivel_bateria"]) == ("Sin rayones.", 90)
    assert acta["vehiculo_codigo"] == "MON-0001"
    assert acta["conformidad"] is None


def test_conformidad_por_la_api_con_la_huella_del_pdf_descargado(cliente_api, entregado) -> None:
    acta_id, cliente = entregado["acta"]["id"], entregado["cliente"]
    pdf = cliente_api.get(f"/actas/{acta_id}/pdf", headers=cliente)
    assert pdf.headers["content-type"] == "application/pdf"

    conformidad = cliente_api.post(f"/actas/{acta_id}/conformidad", headers=cliente).json()[
        "conformidad"
    ]

    # Quien tenga el PDF puede comprobarlo: su SHA-256 es la huella guardada.
    assert conformidad["hash_pdf"] == hashlib.sha256(pdf.content).hexdigest()
    verificacion = cliente_api.get(f"/actas/{acta_id}/verificacion", headers=cliente).json()
    assert verificacion == {
        "con_conformidad": True,
        "coincide": True,
        "hash_pdf": conformidad["hash_pdf"],
    }
    repetida = cliente_api.post(f"/actas/{acta_id}/conformidad", headers=cliente)
    assert repetida.json()["error"]["codigo"] == "conformidad_ya_registrada"


def test_exportaciones_del_mostrador(cliente_api, entregado) -> None:
    operador = entregado["operador"]
    excel = cliente_api.get("/actas/exportacion/excel", headers=operador)
    pdf = cliente_api.get("/actas/exportacion/pdf", params={"tipo": "entrega"}, headers=operador)

    assert excel.status_code == 200
    assert "spreadsheetml" in excel.headers["content-type"]
    assert load_workbook(BytesIO(excel.content)).active.max_row == 4
    assert pdf.content.startswith(b"%PDF")


def test_permisos(cliente_api, crear_usuario, entregado) -> None:
    cliente, operador = entregado["cliente"], entregado["operador"]
    acta_id = entregado["acta"]["id"]
    otro = encabezado_de(crear_usuario(Rol.CLIENTE))

    assert cliente_api.get("/actas", headers=cliente).status_code == 403
    assert cliente_api.get("/actas/exportacion/excel", headers=cliente).status_code == 403
    assert cliente_api.get("/actas/mias", headers=operador).status_code == 403
    assert cliente_api.get(f"/actas/{acta_id}", headers=otro).status_code == 404
    assert cliente_api.get(f"/actas/{acta_id}/pdf", headers=otro).status_code == 404
    assert cliente_api.post(f"/actas/{acta_id}/conformidad", headers=operador).status_code == 403
    assert cliente_api.get("/actas", headers=operador).json()["total"] == 1
